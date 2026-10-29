"""One local, versioned report draft; no map providers, original report writes or agents."""
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import re
import tempfile

from app.schemas.report_draft import DraftChapter, DraftText, ReportDraft

CASE_ID = 'tongji-all7-repaired-20261026'
CASE_IDS = {CASE_ID, 'tongji-all7-small-updated-20261027'}
TITLE = '上海同济大学四平路校区周边15分钟生活圈公共服务设施基础体检与规划研判报告'
EDITABLE = {'findings', 'planning', 'limitations'}

class DraftError(Exception):
    def __init__(self, code, message):
        self.code, self.message = code, message
        super().__init__(message)

def source_hash(report):
    return hashlib.sha256(json.dumps(report.model_dump(mode='json'), ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()).hexdigest()

def locked_facts(report):
    groups = []
    for group in report.groups:
        row = group.model_dump(mode='json')
        coverage = report.blindspot.get('group_coverage', {}).get(group.id, {})
        row.update({k:coverage.get(k) for k in ('covered_ratio','blind_ratio','unknown_ratio')})
        groups.append(row)
    return {'report_id':report.report_id, 'created_at':report.created_at,
            'center':report.center.model_dump(), 'location_label':report.location_label,
            'score':report.score, 'score_provisional':report.score_provisional, 'score_method':report.score_method,
            'groups':groups, 'categories':[c.model_dump(mode='json') for c in report.categories],
            'quality':deepcopy(report.quality), 'warnings':[w.model_dump() for w in report.warnings]}

def directory(report_directory):
    value=os.environ.get('CHECKUP_DRAFT_DIR','').strip()
    return Path(value).expanduser() if value else report_directory.parent/'report_drafts'

def load_draft(root, report, version=None):
    path=root/f'{report.report_id}.json' if version is None else root/'versions'/report.report_id/f'v{version}.json' 
    if path.is_symlink():
        raise DraftError('DRAFT_INVALID','草稿文件不可用，请核对本机文件。')
    if not path.exists():return None
    try:
        draft=ReportDraft.model_validate_json(path.read_text(encoding='utf-8'))
    except (OSError, ValueError):
        raise DraftError('DRAFT_INVALID','草稿文件不可用，未覆盖已有内容。') from None
    if (version is not None and draft.version != version) or draft.report_id != report.report_id or draft.source_hash != source_hash(report) or draft.facts != locked_facts(report):
        raise DraftError('SOURCE_CHANGED','保存报告与草稿依据不一致，请核对报告后再继续。')
    return draft

def _atomic_write(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary=None
    try:
        with tempfile.NamedTemporaryFile(mode='w',encoding='utf-8',dir=path.parent,prefix='.draft-',suffix='.tmp',delete=False) as file:
            temporary=Path(file.name)
            file.write(content); file.flush(); os.fsync(file.fileno())
        os.replace(temporary, path)
    finally:
        if temporary is not None:temporary.unlink(missing_ok=True)

def save_draft(root, draft):
    current=root/f'{draft.report_id}.json'
    versions=root/'versions'/draft.report_id
    # Keep the previous saved version too, including drafts created before version storage.
    if current.exists():
        previous=ReportDraft.model_validate_json(current.read_text(encoding='utf-8'))
        old=versions/f'v{previous.version}.json'
        if not old.exists():_atomic_write(old,previous.model_dump_json(indent=2))
    content=draft.model_dump_json(indent=2)
    _atomic_write(versions/f'v{draft.version}.json',content)
    _atomic_write(current,content)

def percent(value):
    return '未提供' if value is None else f'{value*100:.2f}%'

def count_text(category):
    return str(category['count']) if category['count'] is not None else f"未知（已知{category['known_count']}）"

def initial_draft(root, report):
    initial=load_draft(root/'baselines',report)
    if initial is not None:return initial
    facts=locked_facts(report)
    return ReportDraft(report_id=report.report_id,source_hash=source_hash(report),version=1,
                       title=TITLE,mode='report_facts',notice='依据已保存测算结果整理的完整初稿；尚未进行本次对话的文字修改。',
                       updated_at=report.created_at,facts=facts,chapters=build_chapters(facts))

def build_chapters(facts):
    groups=facts['groups'];categories=facts['categories']
    return [
        DraftChapter(id='scope',title='第一章 分析对象与用途',paragraphs=[
            '本报告面向自然资源和规划主管部门相关业务处室，用于详细规划、城市更新及公共服务设施布局研究的前期基础排查。',
            '分析域为同济大学四平路校区周边指定中心点生成的步行生活圈，不代表完整行政社区、辖区人口普查或法定规划达标结论。'],
            table_headers=['项目','保存结果'],table_rows=[['报告编号',facts['report_id']],['结果时间',facts['created_at']],
            ['分析中心（BD09LL）',f"{facts['center']['lng']}, {facts['center']['lat']}"],
            ['设施配置参考分', '未知' if facts['score'] is None else f"{facts['score']:g}"],
            ['分类范围',f'{len(groups)}大类、{len(categories)}小项']]),
        DraftChapter(id='method',title='第二章 测算口径与解释规则',paragraphs=[
            '分析范围为指定中心的15分钟步行圈，中心采样间距100米。设施服务区以真实POI为起点，以1000米实际步行距离为阈值，通过测距采样及插值估算连续边界。圈外可向圈内提供服务的设施也纳入计算。',
            '覆盖、确认盲区与未知分别保留。地图高亮按既定展示规则包含确认盲区与未知，但未知不代表已经确认无服务。大类内任意小项已有服务，即可排除该处大类共同盲区；大类覆盖充分不代表每个小项齐全。',
            '设施配置参考分用于描述保存数据下的数量配置，不是覆盖率、达标率或规划合规认定。点位分类、设施开放条件、人口需求与容量仍需进一步核查。']),
        DraftChapter(id='findings',title='第三章 设施现状与服务缺口',paragraphs=[
            '下表分别列出各大类的估算覆盖、确认共同盲区和未知比例。应同时查看细分设施结果，避免用大类汇总替代具体设施的服务缺口。'],
            table_headers=['设施大类','覆盖','确认共同盲区','未知'],
            table_rows=[[g['name_zh'],percent(g['covered_ratio']),percent(g['blind_ratio']),percent(g['unknown_ratio'])] for g in groups]),
        DraftChapter(id='planning',title='第四章 规划研判与后续调研建议',paragraphs=[
            '建议将养老与助残设施作为后续专项调研的关注方向，先核查服务对象、开放条件、空间可达性与实际容量，再结合人口结构和现有规划开展布局研究。',
            '小学、药店、菜市场等细分设施应分别核对确认盲区与未知区域，优先补充缺测证据和现场调查。未知不能直接转换成设施缺失，也不能据此直接提出新建学校等选址结论。',
            '可形成设施真实性、开放性、容量与需求的核查清单，供业务处室组织后续调研与方案比选。']),
        DraftChapter(id='limitations',title='第五章 依据与适用边界',paragraphs=[
            '依据为本报告编号所对应的已保存设施分类、步行测距及其质量记录。服务区是采样插值估算，不是对每栋建筑或每个位置的现场逐点核验。',
            '地图POI可能存在漏检、误分类、校园内部楼宇或不向社区开放的设施。尚未获得人口、容量与现场真值的充分证据，现阶段结论仅用于规划前期排查。',
            '本报告生成和文字修订复用已保存结果，不重新测距。原始分析事实保持只读；对本报告的文字调整不会改变地图、分数或测算结果。']),
        DraftChapter(id='appendix',title='附表 设施小项保存结果',paragraphs=[
            '圈内检索数量与设施服务覆盖属于不同指标；圈外设施也可能提供圈内覆盖。未提供比例表示缺少相应测算结果，不按零处理。'],
            table_headers=['设施小项','圈内检索数','覆盖','确认盲区','未知'],
            table_rows=[[c['name_zh'],count_text(c),percent(c['covered_ratio']),percent(c['blind_ratio']),percent(c['unknown_ratio'])] for c in categories]),
    ]

def validate_paragraphs(paragraphs):
    """Editable prose may discuss research tasks; source-owned figures stay in tables."""
    cleaned=[p.strip() for p in paragraphs]
    if any(not p for p in cleaned):
        raise DraftError('FACT_PROTECTION','模型返回空白段落，当前报告未变。')
    text=''.join(cleaned)
    # Do not let free prose introduce measured numbers or definite compliance/absence claims.
    forbidden = [r'[0-9０-９%％]|百分之',
                 r'[零〇一二三四五六七八九十百千万两]+(?:点[零〇一二三四五六七八九]+)?(?:分(?!析|别|类|布|组|为|期|解)|成(?!果|立)|处|家|人|米|公里|平方米|分钟)',
                 r'(?:覆盖率|未知比例|盲区比例|设施配置参考分).{0,8}(?:为零|满分|最高分|百分之)',
                 r'未知.{0,16}(?:已确认无服务|按设施缺失处理|均已确认)',
                 r'(?:所有设施|全部设施|所有小项).{0,8}(?:均已达标|全部达标|均已覆盖)',
                 r'已达到法定规划标准|不存在盲区|没有任何盲区|服务范围不存在',
                 r'(?:建议|应|必须|立即|直接).{0,6}(?:新建|增建|建设).{0,5}学校',
                 r'(?:已完成|已通过|已经完成|均已完成).{0,10}(?:现场|实地|容量|人口).{0,8}(?:核验|核查|调查|复核)|(?:现场|容量|人口).{0,8}(?:调查|核验).{0,3}(?:已完成|完成了)']
    if any(re.search(pattern,text) for pattern in forbidden):
        raise DraftError('FACT_PROTECTION','正文含未经核验的数值或确定性结论，当前报告未变；测算数值由原始表格提供。')
    return cleaned

def produce_draft(report, current, instruction, model, emit, context='', history_root=None):
    facts=locked_facts(report)
    chapters=deepcopy(current.chapters) if current else build_chapters(facts)
    editorial=deepcopy(current.editorial_state) if current else {}
    # Exact chapter-deletion commands do not need a model or a network request.
    command=instruction.strip().rstrip('。！!')
    target=re.fullmatch(r'(?:请|帮我|请帮我)?(?:把|将)?第([一二三四五六七八九十0-9]+)[章张](?:整章)?(?:去掉|删掉|删除|移除)',command)
    target=target or re.fullmatch(r'(?:请|帮我|请帮我)?(?:去掉|删掉|删除|移除)第([一二三四五六七八九十0-9]+)[章张]',command)
    removed=None
    if target and current:
        number=target.group(1)
        if number.isdigit() and 1<=int(number)<=5:number='一二三四五'[int(number)-1]
        removed=next((c for c in chapters if c.title.startswith(f'第{number}章')),None)
        if removed is None:raise DraftError('CHAPTER_NOT_FOUND','当前报告中没有这一章，请核对章节名称后重试。')
        if removed.id not in EDITABLE:raise DraftError('FACT_PROTECTION','分析对象、测算口径和原始结果附表须保留；可删除或改写研判章节。')
        chapters.remove(removed);editorial.pop(removed.id,None)
    restored=None
    restore=re.fullmatch(r'(?:请|帮我|请帮我)?(?:恢复|还原|撤销删除)第([一二三四五六七八九十0-9]+)[章张]',command)
    restore=restore or re.fullmatch(r'(?:请|帮我|请帮我)?(?:把|将)第([一二三四五六七八九十0-9]+)[章张](?:恢复|加回来|加回)',command)
    if restore and current:
        number=restore.group(1)
        if number.isdigit() and 1<=int(number)<=5:number='一二三四五'[int(number)-1]
        original=next((c for c in build_chapters(facts) if c.title.startswith(f'第{number}章')),None)
        if original is None:raise DraftError('CHAPTER_NOT_FOUND','没有找到要恢复的章节。')
        if any(c.id==original.id for c in chapters):raise DraftError('NO_CHANGES','这一章仍在当前报告中，无需恢复。')
        restored=original
        if history_root is not None:
            for version in range(current.version-1,0,-1):
                saved=load_draft(history_root,report,version)
                previous=next((c for c in saved.chapters if c.id==original.id),None) if saved else None
                if previous is not None:restored=deepcopy(previous);break
        chapters.append(restored)
        order={c.id:i for i,c in enumerate(build_chapters(facts))}
        chapters.sort(key=lambda c:order[c.id])
    if instruction and re.search(r'(?:改|设|调整|替换).{0,12}(?:为|成|到)\s*[0-9０-９零一二三四五六七八九十百两]+(?:\.[0-9]+)?\s*(?:分|%|％|成)|(?:未知|盲区|覆盖)(?:率|比例|面积)?.{0,10}(?:改为|改成|设为)\s*[0-9０-９零一二三四五六七八九十百两]+',instruction):
        raise DraftError('FACT_PROTECTION','设施数量、分数及覆盖/盲区/未知属于锁定事实，不能通过对话修改。可修改章节表述、分析重点与建议。')
    if not removed and not restored and model is None and (instruction or current is not None):
        raise DraftError('MODEL_NOT_CONFIGURED','未配置DeepSeek Key，暂不能进行自然语言改写。当前事实版报告可查看和导出；配置后可重试。')
    mode='report_facts'
    notice='未配置DeepSeek Key：已依据保存结果生成事实版模板，未调用模型；自然语言改写暂不可用。'
    if removed:
        mode=current.mode
        notice=f'已按要求删除“{removed.title}”；其余章节保留，原始测算数据与附表保持不变。'
    elif restored:
        mode=current.mode
        notice=f'已恢复“{restored.title}”的最近保存内容；原始测算数据与附表保持不变。'
    elif model is not None:
        emit({'type':'stage','stage':'writing','message':'正在按你的要求撰写报告正文；测算事实与表格保持只读。'})
        payload={'locked_facts':facts,'current_chapters':[c.model_dump() for c in chapters],
                 'instruction':instruction or '生成适合规划主管部门前期调研使用的报告文字。',
                 'conversation_context':context,'editable_ids':sorted(EDITABLE)}
        try:
            raw=model.generate_object(system_prompt=(
                '你为自然资源和规划主管部门撰写公共服务设施前期调研报告。只返回JSON中的sections，每节含id和完整替换段落paragraphs。'
                '请真正重写正文，不能只回答已修改，不能选择固定模板片段。可以扩写、精简、调整重点、顺序、语气，形成具体可读的研判和调研建议。'
                '可编辑findings、planning、limitations；以current_chapters中的id和标题定位，用户说结论默认指planning。只返回用户要求变更的章节。'
                '删除章节时该section仅返回id与remove:true，不提供paragraphs；整章删除允许移除随章只读表格，但不能改动原始事实。已删除章节除非明确要求恢复，不要重新添加。'
                '针对口吻修改，保留当前章节已选主题及内容重点；扩写时补充具体调查对象、核查步骤、形成的成果，不能只加套话。'
                '数值、设施数、覆盖与盲区比例都由服务器在只读表格中呈现。正文不写任何阿拉伯数字、百分号或新的测算数值，也不将数值改用汉字复述。'
                '只能依据locked_facts解读已保存结果；对话上下文仅帮助理解用户意图，不是事实来源。不得声称完成了现场调查、人口容量核验或法定达标。'
                '未知不等于缺失，大类覆盖不等于每个小项齐全。建议限于核查、调研和条件明确后的方案比选，不直接提出新建学校等选址结论。'
                '不得修改章节标题、固定口径章节、附表、原始数据；不要输出Markdown标题或表格；分点内容用独立段落，不用数字编号。'),
                user_message=json.dumps(payload,ensure_ascii=False),response_schema=DraftText.model_json_schema())
        except Exception:
            raise DraftError('MODEL_FAILED','DeepSeek生成失败，当前报告未变。可在同一报告上重试，无需重新体检。') from None
        try:
            replacement=DraftText.model_validate(raw)
        except ValueError:
            raise DraftError('FACT_PROTECTION','模型返回了非约定章节或字段，未替换当前报告；可重试。') from None
        seen=set()
        for section in replacement.sections:
            if section.id in seen:raise DraftError('FACT_PROTECTION','模型返回重复章节，当前报告未变。')
            seen.add(section.id)
            chapter=next((c for c in chapters if c.id==section.id),None)
            if chapter is None:raise DraftError('CHAPTER_NOT_FOUND','请求涉及已删除章节，当前报告未变；请核对章节名称。')
            if section.remove:
                chapters.remove(chapter)
            elif section.paragraphs != chapter.paragraphs:
                chapter.paragraphs=validate_paragraphs(section.paragraphs)
            # Invalidate legacy block selections for rewritten chapters: prose is now authoritative.
            editorial.pop(section.id,None)
        if instruction and current and chapters==current.chapters:
            raise DraftError('NO_CHANGES','本次没有产生正文变更，已保留当前报告。请说明需要补充、删除或改写的具体内容。')
        mode='model';notice='DeepSeek已根据当前保存结果撰写本版研判正文；测算事实与表格保持只读，建议供前期调研复核。'
    emit({'type':'stage','stage':'validation','message':'正在校验章节、锁定事实与当前版本。'})
    for chapter in chapters:
        emit({'type':'chapter','chapter':chapter.model_dump()})
    return ReportDraft(report_id=report.report_id,source_hash=source_hash(report),version=(current.version if current else 0)+1,
                       title=TITLE,mode=mode,notice=notice,updated_at=datetime.now(timezone.utc).isoformat(),facts=facts,chapters=chapters,editorial_state=editorial)

def export_docx(draft):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document=Document();section=document.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.8)
    section.left_margin=section.right_margin=Inches(.75)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=document.styles[name];style.font.name='Songti SC';style.font.color.rgb=RGBColor(0,0,0)
        fonts=style.element.get_or_add_rPr().rFonts
        for key in list(fonts.attrib):
            if 'theme' in key.lower():del fonts.attrib[key]
        fonts.set(qn('w:eastAsia'),'Songti SC')
        borders=style.element.find(qn('w:pPr'))
        if borders is not None:
            border=borders.find(qn('w:pBdr'))
            if border is not None:borders.remove(border)
    normal=document.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.25
    document.styles['Title'].font.size=Pt(20)
    document.styles['Heading 1'].font.size=Pt(14)
    document.add_heading(draft.title,0)
    document.add_paragraph(f'依据报告：{draft.report_id}')
    document.add_paragraph(draft.notice)
    for chapter in draft.chapters:
        document.add_heading(chapter.title,1)
        for text in chapter.paragraphs:document.add_paragraph(text)
        if chapter.table_headers:
            table=document.add_table(rows=1, cols=len(chapter.table_headers));table.style='Table Grid'
            widths=([2.0,5.0] if len(chapter.table_headers)==2 else
                    [2.5,1.5,1.5,1.5] if len(chapter.table_headers)==4 else [2.3,1.1,1.2,1.2,1.2])
            table.autofit=False
            for column,width in zip(table.columns,widths):column.width=Inches(width)
            for cell,width,label in zip(table.rows[0].cells,widths,chapter.table_headers):
                cell.width=Inches(width);cell.text=label
                for run in cell.paragraphs[0].runs:run.bold=True
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDEDED');cell._tc.get_or_add_tcPr().append(shade)
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in chapter.table_rows:
                cells=table.add_row().cells
                for cell,width,value in zip(cells,widths,row):cell.width=Inches(width);cell.text=value
            for row in table.rows:
                no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
            document.add_paragraph()
    stream=BytesIO();document.save(stream);return stream.getvalue()
